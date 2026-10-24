"""Builds the test-data template for Turbopump: one data sheet that exports to a single CSV."""
import csv, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.comments import Comment

OUT = sys.argv[1]
FONT = "Arial"
title = Font(name=FONT, size=14, bold=True)
bold = Font(name=FONT, bold=True)
plain = Font(name=FONT)
note = Font(name=FONT, italic=True, color="666666")
inp = Font(name=FONT, color="0000FF")
fill = PatternFill("solid", fgColor="FFF2B2")
head_fill = PatternFill("solid", fgColor="DDE5F0")
thin = Side(style="thin", color="AAAAAA")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

# (key, example value, unit, description, required)
PUMP = [
    ("nombre", "Ejemplo 65-160", "", "Nombre o referencia de la bomba", True),
    ("caudal", 100, "m3/h", "Caudal del punto de diseño (o del óptimo medido)", True),
    ("altura", 32, "m", "Altura del punto de diseño", True),
    ("velocidad", 2900, "rpm", "Velocidad nominal", True),
    ("fluido", "Water", "", "Fluido (nombre de CoolProp; agua: Water)", True),
    ("temperatura", 20, "°C", "Temperatura del fluido en el ensayo", True),
    ("presion_entrada", 101.325, "kPa", "Presión total absoluta en la aspiración (para el NPSH disponible)", False),
    ("d2", 172, "mm", "Diámetro exterior del rodete d₂ (el ensayado, recortado si lo está)", True),
    ("b2", 17.5, "mm", "Ancho de salida b₂", True),
    ("z", 7, "", "Número de álabes", True),
    ("beta2", 22, "°", "Ángulo de salida del álabe β₂B, desde la dirección circunferencial", True),
    ("d1", 77, "mm", "Diámetro del ojo d₁ (en la cubierta)", True),
    ("d_cubo", 19, "mm", "Diámetro del cubo en el ojo", True),
    ("espesor", 3, "mm", "Espesor del álabe", False),
    ("aspiracion", "simple", "", "simple o doble", True),
    ("d_anillo", "", "mm", "Diámetro del anillo de desgaste delantero (opcional)", False),
    ("holgura_anillo", "", "mm", "Holgura radial del anillo (opcional)", False),
    ("voluta", "simple", "", "simple o doble", True),
    ("dz", 186, "mm", "Diámetro del corte de agua d_z", True),
    ("b3", 30, "mm", "Ancho de entrada de la voluta b₃", True),
    ("garganta", 3100, "mm2", "Área de la garganta (si es doble, la suma de las dos)", False),
    ("nervio", "", "°", "Ángulo del nervio de la voluta doble, desde el corte de agua (opcional)", False),
    ("velocidad_ensayo", 2900, "rpm", "Velocidad a la que se midió (si difiere, se corrige por semejanza)", True),
    ("norma", "ISO 9906 grado 2B", "", "Norma y grado del ensayo", False),
    ("fecha", "2026-10-22", "", "Fecha del ensayo", False),
    ("banco", "Banco 1", "", "Banco o laboratorio", False),
]
TEST_HEAD = ["caudal", "altura", "potencia", "rendimiento", "npsh3"]
TEST_UNITS = ["m3/h", "m", "kW", "%", "m"]
TEST_DESC = ["Q medido", "H medida", "P en el eje", "η total (opcional)", "NPSH₃ (opcional)"]
EXAMPLE = [[0, 39.5, 4.1, 0, ""], [50, 36.6, 7.9, 63.0, 1.9], [100, 32.1, 10.8, 80.9, 3.3], [130, 27.4, 12.6, 76.9, 4.9]]

wb = Workbook()
ws = wb.active
ws.title = "Datos"
ws["A1"] = "Ensayo de bomba para Turbopump"; ws["A1"].font = title
ws["A2"] = "Rellena las celdas amarillas. Los valores azules son un EJEMPLO: sustitúyelos. Para Turbopump, exporta esta hoja con «Guardar como → CSV UTF-8»."
ws["A2"].font = note
ws["A4"] = "[BOMBA]"; ws["A4"].font = bold
for c, h in enumerate(["clave", "valor", "unidad", "descripción"], 1):
    cell = ws.cell(row=5, column=c, value=h); cell.font = bold; cell.fill = head_fill; cell.border = box
r = 6
for key, value, unit, desc, req in PUMP:
    ws.cell(row=r, column=1, value=key).font = plain
    v = ws.cell(row=r, column=2, value=value); v.font = inp; v.fill = fill; v.border = box
    u = ws.cell(row=r, column=3, value=unit); u.font = inp; u.fill = fill; u.border = box
    d = ws.cell(row=r, column=4, value=desc if req or "(opcional)" in desc else desc + " (opcional)")
    d.font = plain if req else note
    for col in (1, 4):
        ws.cell(row=r, column=col).border = box
    r += 1
r += 1
ws.cell(row=r, column=1, value="[ENSAYO]").font = bold
r += 1
for c, h in enumerate(TEST_HEAD, 1):
    cell = ws.cell(row=r, column=c, value=h); cell.font = bold; cell.fill = head_fill; cell.border = box
header_row = r
r += 1
for c, u in enumerate(TEST_UNITS, 1):
    cell = ws.cell(row=r, column=c, value=u); cell.font = inp; cell.fill = fill; cell.border = box
units_row = r
r += 1
for c, d in enumerate(TEST_DESC, 1):
    ws.cell(row=r, column=c, value=d).font = note
desc_row = r
r += 1
first_point = r
for row in EXAMPLE:
    for c, v in enumerate(row, 1):
        cell = ws.cell(row=r, column=c, value=v); cell.font = inp; cell.fill = fill; cell.border = box
    r += 1
for _ in range(16):
    for c in range(1, 6):
        cell = ws.cell(row=r, column=c); cell.fill = fill; cell.border = box
    r += 1
ws.cell(row=header_row, column=1).comment = Comment("Una fila por punto medido, de válvula cerrada a sobrecarga. Puntos crudos del banco mejor que una curva suavizada.", "Turbopump")
ws.cell(row=units_row, column=1).comment = Comment("Unidades editables: m3/h, l/s, gpm; m, ft; kW, hp; % o fracción.", "Turbopump")
ws.column_dimensions["A"].width = 18
ws.column_dimensions["B"].width = 20
ws.column_dimensions["C"].width = 10
ws.column_dimensions["D"].width = 80
ws.column_dimensions["E"].width = 12
ws.freeze_panes = "A6"

ins = wb.create_sheet("Instrucciones")
lines = [
    ("Cómo rellenar la plantilla", title),
    ("", plain),
    ("1. En la hoja «Datos», escribe solo en las celdas amarillas. Lo azul es un ejemplo: sustitúyelo o bórralo.", plain),
    ("2. Bloque [BOMBA]: un dato por fila. No cambies la columna «clave»: Turbopump la usa para reconocer cada dato.", plain),
    ("3. Las unidades se pueden cambiar (mm o in, m3/h o gpm, m o ft, kW o hp). Turbopump las convierte.", plain),
    ("4. Los datos con la descripción en gris cursiva son opcionales: deja el valor vacío si no lo tienes.", plain),
    ("5. Bloque [ENSAYO]: una fila por punto medido, de válvula cerrada a sobrecarga. Rendimiento y NPSH₃ son opcionales.", plain),
    ("6. Si la velocidad del ensayo no es la nominal, indícala en «velocidad_ensayo»; Turbopump corrige los puntos por semejanza.", plain),
    ("7. Para Turbopump, exporta solo la hoja «Datos»: Archivo → Guardar como → CSV UTF-8. Sirven la coma o el punto y coma como separador, y la coma o el punto decimal.", plain),
    ("", plain),
    ("Confidencialidad", bold),
    ("Los ensayos y la geometría de los rodetes son datos de la empresa. Confirma que puedes usarlos antes de cargarlos.", plain),
    ("Turbopump los compara con su predicción; no se publican ni se usan para ajustar coeficientes a una bomba concreta.", plain),
    ("Los resultados son predicciones de diseño preliminar: no certifican conformidad con ISO 9906 ni con ninguna norma.", plain),
]
for i, (text, f) in enumerate(lines, 1):
    c = ins.cell(row=i, column=1, value=text); c.font = f; c.alignment = Alignment(wrap_text=False)
ins.column_dimensions["A"].width = 120
wb.save(OUT)

# The CSV that "Save as CSV" of the data sheet gives, for the loader's tests.
with open(OUT.replace(".xlsx", "-ejemplo.csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    for row in ws.iter_rows(values_only=True):
        w.writerow(["" if v is None else v for v in row])
print("ok", first_point)
